import os
import ast
import asyncio
import nest_asyncio
import streamlit as st
import pdfplumber
import docx
import pandas as pd
import matplotlib.pyplot as plt
import time
import random
import uuid
from datetime import datetime
from io import BytesIO, StringIO
from fpdf import FPDF
from langchain_groq import ChatGroq
from langchain.schema import HumanMessage, SystemMessage
from langchain.memory import ConversationBufferMemory

# Allow asyncio to run in Streamlit
nest_asyncio.apply()

###############################################################################
# CONFIGURATION & SETUP
###############################################################################
st.set_page_config(
    page_title="EduMind AI: MCQ Generator & Chatbot",
    page_icon="🧠",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Initialize session state variables
if "initialized" not in st.session_state:
    st.session_state.initialized = True
    st.session_state.chat_history = []
    st.session_state.all_chats = {}
    st.session_state.current_chat_id = str(uuid.uuid4())
    st.session_state.all_chats[st.session_state.current_chat_id] = {
        "title": "New Chat",
        "messages": [],
        "created_at": datetime.now()
    }
    st.session_state.mcq_history = {}
    st.session_state.memory = ConversationBufferMemory(memory_key="chat_history", return_messages=True)
    st.session_state.show_welcome = True
    st.session_state.default_num_questions = 10
    st.session_state.default_difficulty = "Medium"
    st.session_state.mcq_topic = ""
    st.session_state.mcq_questions = []
    st.session_state.mcq_current_question = 0
    st.session_state.mcq_answers = []
    st.session_state.mcq_test_completed = False
    st.session_state.mcq_test_id = None
    st.session_state.file_text = ""
    st.session_state.topics = []

# Set your Groq API key
GROQ_API_KEY = os.getenv("GROQ_API_KEY", "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")
os.environ["GROQ_API_KEY"] = GROQ_API_KEY

# Initialize chat model
chat = ChatGroq(temperature=0.7, model_name="llama3-70b-8192", groq_api_key=GROQ_API_KEY)
def load_css():
    return """
    <style>
    /* Global Styles */
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700&display=swap');
    
    :root {
        --primary-color: #4A56E2;
        --secondary-color: #6E7CE6;
        --accent-color: #8B5CF6;
        --background-dark: #111827;
        --card-bg: #1F2937;
        --sidebar-bg: #111827;
        --text-color: #E5E7EB;
        --text-secondary: #9CA3AF;
        --success-color: #10B981;
        --warning-color: #F59E0B;
        --error-color: #EF4444;
        --border-color: #374151;
    }
    
    body {
        font-family: 'Inter', sans-serif;
        background-color: var(--background-dark);
        color: var(--text-color);
    }
    
    /* Chat message styling */
    .chat-message {
        padding: 1rem;
        border-radius: 0.7rem;
        margin-bottom: 1rem;
        display: flex;
        box-shadow: 0 2px 6px rgba(0, 0, 0, 0.1);
    }
    
    .chat-message.user {
        background-color: var(--primary-color);
        border-bottom-right-radius: 0.2rem;
        margin-left: 2rem;
    }
    
    .chat-message.assistant {
        background-color: var(--card-bg);
        border-bottom-left-radius: 0.2rem;
        margin-right: 2rem;
    }
    
    /* Additional styling... */
    </style>
    """
def render_sidebar():
    with st.sidebar:
        st.sidebar.image("https://img.icons8.com/clouds/100/000000/brain.png", width=80)
        st.sidebar.title("EduMind AI")
        
        # New Chat button
        if st.sidebar.button("New Chat", key="new_chat"):
            st.session_state.chat_history = []
            st.session_state.current_chat_id = str(uuid.uuid4())
            st.session_state.all_chats[st.session_state.current_chat_id] = {
                "title": "New Chat",
                "messages": [],
                "created_at": datetime.now()
            }
            st.session_state.memory.clear()
            st.session_state.show_welcome = True
            rerun_app()
        
        # Mode selection
        st.sidebar.markdown("### Mode")
        mode = st.sidebar.radio(
            "Select Mode:",
            ["MCQ Generator", "Educational Chatbot"],
            key="mode_selection"
        )
        
        # Chat history section
        st.sidebar.markdown("### Chat History")
        if len(st.session_state.all_chats) > 0:
            for chat_id, chat_data in sorted(
                st.session_state.all_chats.items(), 
                key=lambda x: x[1]["created_at"], 
                reverse=True
            ):
                # Display chat history entries
                # Additional implementation details...
def generate_mcqs(content, difficulty, num_questions):
    """Generate MCQs synchronously by wrapping the async function."""
    loop = asyncio.new_event_loop()
    asyncio.set_event_loop(loop)
    result = loop.run_until_complete(async_generate_mcqs(content, difficulty, num_questions))
    loop.close()
    
    try:
        # Extract and validate the Python list of questions
        if "
            code_block = result.split("python")[1].split("
            data = ast.literal_eval(code_block)
        else:
            data = ast.literal_eval(result)
        
        if not isinstance(data, list):
            raise ValueError("Output is not a list")
        
        # Validate question format
        validated_questions = []
        for q in data:
            if not isinstance(q, list):
                continue
                
            # Check if it's a 4-option MCQ or True/False
            if len(q) == 7:  # 4-option MCQ
                validated_questions.append(q)
            elif len(q) == 5:  # True/False
                validated_questions.append(q)
        
        return validated_questions[:num_questions]
    except Exception as e:
        st.error(f"Error parsing MCQ output: {e}")
        return []
def render_mcq_generator():
    st.title("MCQ Generator")
    
    # MCQ generator form (when no test is in progress)
    st.write("Create a new MCQ test by entering a topic, uploading a document, or pasting text.")
    
    method_tabs = st.tabs(["Enter Topic", "Upload Document", "Paste Text"])
    
    with method_tabs:
        topic = st.text_input("Enter a topic for your MCQ test:", 
                             placeholder="e.g., Python Programming, World War II, Organic Chemistry")
        if topic:
            st.session_state.mcq_topic = topic
    
    with method_tabs[1]:
        uploaded_file = st.file_uploader("Upload a document:", type=["pdf", "docx", "txt"])
        if uploaded_file:
            with st.spinner("Extracting text..."):
                st.session_state.file_text = extract_text(uploaded_file, uploaded_file.name)
                st.success(f"Successfully extracted {len(st.session_state.file_text)} characters.")
            
            # Show a preview of the extracted text
            with st.expander("Preview extracted text"):
                st.text(st.session_state.file_text[:1000] + "..." if len(st.session_state.file_text) > 1000 else st.session_state.file_text)
    
    with method_tabs:
        text_area = st.text_area("Paste your text content:", height=200, 
                               placeholder="Paste educational content here...")
        if text_area:
            st.session_state.file_text = text_area
            st.write(f"Character count: {len(text_area)}")
def suggest_topics(text):
    """Use LLM to suggest topics from text content."""
    system_prompt = """
    You are an educational content analyzer. Given a text, identify 5-7 key topics that would make good 
    MCQ test subjects. Return only a Python list of topic strings without any additional text.
    """
    
    try:
        messages = [
            SystemMessage(content=system_prompt),
            HumanMessage(content=f"Analyze this text and suggest MCQ topics:\n\n{text[:30000]}")
        ]
        response = chat.invoke(messages)
        topics = ast.literal_eval(response.content)
        return topics[:7]  # Limit to 7 topics
    except Exception as e:
        st.error(f"Error suggesting topics: {e}")
        return ["General Knowledge"]
# Display progress bar and counter
st.markdown(f"""
<div style="display: flex; justify-content: space-between; align-items: center; margin-bottom: 0.5rem;">
    <span>Question {idx + 1} of {total}</span>
    <span>{progress_pct:.0f}% Complete</span>
</div>
<div class="progress-container">
    <div class="progress-bar" style="width: {progress_pct}%;"></div>
</div>
""", unsafe_allow_html=True)

# Timer display if enabled
if "test_start_time" in st.session_state and "timer_duration" in st.session_state:
    elapsed = time.time() - st.session_state.test_start_time
    remaining = max(0, st.session_state.timer_duration - elapsed)
    
    if remaining <= 0:
        st.warning("Time's up! Please complete your answers.")
    else:
        st.markdown(f"""
        <div class="timer">⏰ Time Remaining: {format_time(remaining)}</div>
        """, unsafe_allow_html=True)

# Display question and options
st.markdown(f"<h3 class='mcq-question'>{mcq}</h3>", unsafe_allow_html=True)

# Option selection interface
for letter, text in options:
    key = f"q_{idx}_{letter}"
    col1, col2 = st.columns([1])
    
    with col1:
        if st.checkbox("", key=key, value=st.session_state.mcq_answers[idx] == letter):
            selected_option = letter
            st.session_state.mcq_answers[idx] = letter
    
    with col2:
        st.markdown(f"<div class='mcq-option {'selected' if st.session_state.mcq_answers[idx] == letter else ''}'><strong>{letter}.</strong> {text}</div>", unsafe_allow_html=True)
def create_interactive_mcq_results(mcqs, user_answers, topic, time_taken=None):
    """Generate HTML-formatted interactive MCQ results for display in Streamlit."""
    score = 0
    html_content = f"""
    <div class="card" style="padding: 1.5rem; margin-bottom: 2rem;">
        <h2 style="margin-top: 0; color: #4A56E2;">MCQ Test Results</h2>
        <p><strong>Topic:</strong> {topic}</p>
        <p><strong>Date:</strong> {datetime.now().strftime('%B %d, %Y, %H:%M')}</p>
        <p><strong>Time Taken:</strong> {time_taken if time_taken else "N/A"}</p>
    </div>
    """
    
    for i, (mcq, ans) in enumerate(zip(mcqs, user_answers)):
        # Result calculation and formatting logic
        # Generate HTML for each question with color-coded feedback
        
    # Add final score
    percentage = round((score / len(mcqs)) * 100)
    html_content += f"""
        <div class="final-score" style="margin-top: 2rem; text-align: center;">
            <h2 style="color: #E5E7EB;">Final Score: {score} / {len(mcqs)} ({percentage}%)</h2>
        </div>
    """
    
    return html_content
def create_pdf_summary(mcqs, user_answers, topic, time_taken=None):
    """Generate a PDF with MCQ test results and explanations."""
    score = 0
    pdf = FPDF()
    pdf.add_page()
    
    # Set up PDF styling
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, "MCQ Test Results", ln=True, align="C")
    
    # Add metadata
    pdf.set_font("Arial", "B", 12)
    pdf.cell(0, 10, f"Topic: {topic}", ln=True)
    pdf.cell(0, 10, f"Date: {datetime.now().strftime('%B %d, %Y, %H:%M')}", ln=True)
    if time_taken:
        pdf.cell(0, 10, f"Time Taken: {time_taken}", ln=True)
    
    # Process each question and create formatted PDF content
    
    return pdf.output(dest="S").encode("latin1")
def render_chatbot():
    st.title("Educational Chatbot")
    
    # Welcome screen for new chats
    if st.session_state.show_welcome:
        st.markdown("""
        <div class="welcome-card">
            <h1 class="welcome-title">Welcome to EduMind AI! 👋</h1>
            <p class="welcome-subtitle">I'm your educational assistant, ready to help you learn and understand any topic.</p>
        </div>
        """, unsafe_allow_html=True)
        
        # Feature highlights
        col1, col2 = st.columns(2)
        with col1:
            st.markdown("""
            <div class="feature-card">
                <h3 class="feature-title">✨ Ask Anything</h3>
                <p>Get detailed explanations on complex topics from math to literature.</p>
            </div>
            """, unsafe_allow_html=True)
            
            # Additional feature cards...
        
        # Example questions
        st.markdown("### Try asking me:")
        example_col1, example_col2, example_col3 = st.columns(3)
        
        example_questions = [
            "Explain the theory of relativity simply",
            "What are the key themes in Hamlet?",
            "How do I solve quadratic equations?",
            # Additional examples...
        ]
# Display chat messages
if st.session_state.chat_history:
    for i, message in enumerate(st.session_state.chat_history):
        if message["role"] == "user":
            st.markdown(f"""
            <div class="chat-message user">
                <div class="chat-avatar">👤</div>
                <div class="chat-content">{message["content"]}</div>
            </div>
            """, unsafe_allow_html=True)
        else:
            st.markdown(f"""
            <div class="chat-message assistant">
                <div class="chat-avatar">🤖</div>
                <div class="chat-content">{message["content"]}</div>
            </div>
            """, unsafe_allow_html=True)
def generate_performance_chart(mcq_history):
    """Generate a performance chart based on MCQ history."""
    if not mcq_history:
        return None
    
    # Extract data for visualization
    dates = []
    scores = []
    topics = []
    
    for test_id, test_data in mcq_history.items():
        if 'date' in test_data and 'score' in test_data and 'total' in test_data:
            dates.append(test_data['date'])
            scores.append((test_data['score'] / test_data['total']) * 100)
            topics.append(test_data['topic'])
    
    # Create visualization
    df = pd.DataFrame({
        'Date': dates,
        'Score (%)': scores,
        'Topic': topics
    })
    
    # Generate and return matplotlib figure
    # ...
def main():
    # Load custom CSS
    st.markdown(load_css(), unsafe_allow_html=True)
    
    # Render sidebar
    render_sidebar()
    
    # Main content based on selected mode
    if st.session_state.get("mode_selection") == "MCQ Generator":
        render_mcq_generator()
    else:
        render_chatbot()

if __name__ == "__main__":
    main()
